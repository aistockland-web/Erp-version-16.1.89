#!/usr/bin/env python3
"""Build canonical Stockland ERP 16.1.91 from the checksum-verified 16.1.90 installer."""
import copy, hashlib, json, re, zipfile
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
BASE = REPO / "STOCKLAND_ERP_16.1.90_CANONICAL_WORDPRESS_UPDATE.zip"
BASE_SHA = "6a7619c834fb37ef69024e68f15ae524efae39b5586b01441ab47270842588e5"
OUT = REPO / "STOCKLAND_ERP_16.1.91_CANONICAL_WORDPRESS_UPDATE.zip"
CSS = r"""/* 16.1.91 — consistent notification typography for every authorized role. */
.sl-overlay .sl-alert-drawer{font-family:var(--sl-font);font-size:14px}
.sl-overlay .sl-alert-drawer .sl-drawer-head h2{font-size:clamp(20px,2.2vw,24px);font-weight:800;line-height:1.25}
.sl-overlay .sl-alert-drawer .sl-drawer-head p{font-size:13px;line-height:1.55}
.sl-overlay .sl-alert-drawer .sl-alert-summary{font-size:12px;line-height:1.5}
.sl-overlay .sl-alert-drawer .sl-notification-card{gap:12px;padding:15px 16px;border-left:3px solid var(--sl-accent,var(--sl-cobalt))}
.sl-overlay .sl-alert-drawer .sl-notification-card.action{border-left-color:var(--sl-cobalt)}
.sl-overlay .sl-alert-drawer .sl-notification-card.success{border-left-color:var(--sl-emerald)}
.sl-overlay .sl-alert-drawer .sl-notification-card.warning{border-left-color:var(--sl-amber)}
.sl-overlay .sl-alert-drawer .sl-notification-card.danger{border-left-color:var(--sl-red)}
.sl-overlay .sl-alert-drawer .sl-notification-card.info{border-left-color:var(--sl-teal)}
.sl-overlay .sl-alert-drawer .sl-alert-main strong{font-size:14px;font-weight:800;line-height:1.45;letter-spacing:-.01em}
.sl-overlay .sl-alert-drawer .sl-alert-main .sl-sub{margin-top:5px;font-size:13px;line-height:1.6}
.sl-overlay .sl-alert-drawer .sl-alert-main small{margin-top:8px;font-size:11px;line-height:1.5;font-variant-numeric:tabular-nums}
.sl-overlay .sl-alert-drawer .sl-notification-actions .sl-btn{min-height:40px;padding:9px 12px;font-size:12px;font-weight:750;line-height:1.35}
.sl-overlay .sl-alert-drawer .sl-notification-actions .sl-btn:disabled{cursor:wait;opacity:.66}
.sl-toast{font-size:13px;font-weight:700;line-height:1.5}
@media (max-width:767.98px){.sl-overlay .sl-alert-drawer .sl-notification-card{gap:10px;padding:13px 12px}.sl-overlay .sl-alert-drawer .sl-alert-main strong{font-size:14px}.sl-overlay .sl-alert-drawer .sl-alert-main .sl-sub{font-size:13px}.sl-overlay .sl-alert-drawer .sl-notification-actions .sl-btn{flex:1 1 135px;min-height:44px;font-size:12px}}
@media (prefers-reduced-motion:reduce){.sl-overlay .sl-alert-drawer .sl-notification-card{transition:none}}
"""
def need(ok, msg):
    if not ok:
        raise SystemExit("STOP: " + msg)

def css_normalize(value):
    value = re.sub(r"/\*[\s\S]*?\*/", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    value = re.sub(r"\s*([{}:;,])\s*", r"\1", value)
    return value.replace(";}", "}")

def canonicalize_xlsx(path):
    normalized = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as src:
        entries = [(name, src.read(name), src.getinfo(name).external_attr) for name in sorted(src.namelist())]
    with zipfile.ZipFile(normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as dst:
        for name, data, attrs in entries:
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = attrs
            info.create_system = 3
            dst.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    path.write_bytes(normalized.read_bytes())
    normalized.unlink()

def main():
    need(BASE.is_file(), "16.1.90 base ZIP missing")
    need(hashlib.sha256(BASE.read_bytes()).hexdigest() == BASE_SHA, "base SHA mismatch")
    with zipfile.ZipFile(BASE) as z:
        need(z.testzip() is None, "base ZIP CRC failed")
        names = z.namelist()
        need(len(names) == 74 and len(set(names)) == 74 and all(n.startswith("stockland-erp-app/") for n in names), "unexpected base ZIP inventory")
        with TemporaryDirectory(prefix="stockland-1691-") as td:
            z.extractall(td)
            app = Path(td) / "stockland-erp-app"
            plugin_path = app / "stockland-inventory.php"
            s = plugin_path.read_text()
            need("Version:           16.1.90" in s, "plugin release marker missing")
            s = s.replace("Version:           16.1.90", "Version:           16.1.91", 1)
            s = s.replace("define( 'STOCKLAND_VERSION', '16.1.90' );", "define( 'STOCKLAND_VERSION', '16.1.91' );", 1)
            plugin_path.write_text(s)
            front_path = app / "includes/class-stockland-frontend.php"
            s = front_path.read_text()
            need("const ROUTE_VERSION      = '16.1.90';" in s, "route/cache version missing")
            front_path.write_text(s.replace("const ROUTE_VERSION      = '16.1.90';", "const ROUTE_VERSION      = '16.1.91';", 1))

            mp = app / "assets/css/modules.json"
            md = json.loads(mp.read_text())
            need(len(md["modules"]) == 6, "expected six ordered CSS modules")
            for name in md["modules"]:
                p = app / "assets/css" / name
                t = p.read_text()
                if "Stockland ERP 16.1.89 •" in t:
                    t = t.replace("Stockland ERP 16.1.89 •", "Stockland ERP 16.1.91 •", 1)
                if name == "06-operational-guards.css":
                    need("consistent notification typography for every authorized role" not in t, "notification release rules already present in baseline")
                    t = t.rstrip() + "\n\n" + CSS + "\n"
                p.write_text(t)
            chunks = [(app / "assets/css" / n).read_text().strip() for n in md["modules"]]
            (app / "assets/app.css").write_text("\n\n".join(chunks) + "\n")
            min_path = app / "assets/app.min.css"
            min_path.write_text(min_path.read_text().rstrip() + css_normalize(CSS) + "\n")
            md["release"] = "16.1.91"
            md["description"] = "16.1.91 notification typography/severity cues and editable workbook date/history guidance; business logic and role action hooks unchanged."
            mp.write_text(json.dumps(md, ensure_ascii=False, indent=2) + "\n")

            physical = app / "templates/Stockland_Physical_Count_Template.xlsx"
            wb = load_workbook(physical)
            need(not wb["PHYSICAL COUNT"].protection.sheet and not wb["INSTRUCTIONS"].protection.sheet, "physical count workbook must remain unprotected")
            wb["PHYSICAL COUNT"]["A2"] = "IMPORTANT: Count every Product Master SKU. If a variant is not physically present, enter Quantity Found = 0. Fill Brand, Size and Color where applicable. One location per workbook. Enter Count Date as YYYY-MM-DD."
            wb["PHYSICAL COUNT"].row_dimensions[2].height = 40
            wb.properties.modified = datetime(2026, 1, 1, 0, 0, 0)
            ins = wb["INSTRUCTIONS"]
            for row, value in [
                (12, "11. Count Date is the date the physical count was actually taken (YYYY-MM-DD). It does not rewrite a transaction date or SKU effective date."),
                (13, "12. This workbook is unprotected for entry. Do not edit/unlock posted stock-ledger history or historical SKU effective dates. Correct differences through the approved recount/adjustment workflow with a reason and required approval.")
            ]:
                c = ins.cell(row, 1); c.value = value
                c._style = copy.copy(ins["A11"]._style); c.font = copy.copy(ins["A11"].font)
                c.alignment = copy.copy(ins["A11"].alignment); c.fill = copy.copy(ins["A11"].fill)
                c.border = copy.copy(ins["A11"].border)
                ins.row_dimensions[row].height = 36 if row == 12 else 42
            wb.save(physical)
            canonicalize_xlsx(physical)

            master = app / "templates/Stockland_Master_Import_Template.xlsx"
            wb = load_workbook(master)
            need(all(not sheet.protection.sheet for sheet in wb.worksheets), "master workbook must remain unprotected")
            wb["PRODUCT MASTER"]["A2"] = "ENTER STOCKLAND PRODUCT DATA ONLY. Primary SKU must begin with STK-. Product Master never creates physical stock; use Warehouse Receiving for quantity. Do not overwrite posted stock, SKU history or effective dates."
            wb.properties.modified = datetime(2026, 1, 1, 0, 0, 0)
            rules = wb["IMPORT RULES"]
            rules["A9"] = "SKU history / dates"
            rules["B9"] = "Historical effective dates and posted stock are system-controlled; use the authorized SKU-change or stock workflow. Do not edit history directly in Excel."
            for dst, src in [("A9", "A8"), ("B9", "B8")]:
                c = rules[dst]; c._style = copy.copy(rules[src]._style); c.font = copy.copy(rules[src].font)
                c.fill = copy.copy(rules[src].fill); c.border = copy.copy(rules[src].border)
                c.alignment = copy.copy(rules[src].alignment)
            rules.row_dimensions[9].height = 36
            wb.save(master)
            canonicalize_xlsx(master)

            readme = app / "README.md"
            s = readme.read_text()
            need(s.startswith("# Stockland Ethiopia ERP 16.1.90"), "README baseline mismatch")
            s = s.replace("# Stockland Ethiopia ERP 16.1.90", "# Stockland Ethiopia ERP 16.1.91", 1)
            intro = """## Notification readability and inventory date-entry clarity — 16.1.91

- Uses shared notification drawer typography for all authorized ERP roles: larger titles/messages/metadata, clear action/warning/danger/info accent rails, readable toast text, and responsive touch targets. Existing notification data, authorization, action IDs, and handlers are unchanged.
- Both Excel workbooks remain unprotected for normal entry. Count Date is the date the physical count was taken; it does not rewrite a past sale/receipt or historical SKU effective date.
- Product Master import remains quantity-free. Posted stock-ledger history and SKU effective dates must not be edited directly in Excel; use existing approved recount/adjustment or SKU-change workflows with reason and audit trail.
- Plugin/front-end route/cache identity moves to 16.1.91. No JavaScript, REST business behavior, role/capability, payroll, schema, inventory ledger, or business-record code was changed.

**Acceptance boundary:** the candidate has not been installed on live WordPress; exact-version GitHub checks and live production acceptance are tracked separately.

"""
            anchor = "## Owner-only wholesale surplus protection — 16.1.87"
            need(anchor in s, "README insert anchor missing")
            readme.write_text(s.replace(anchor, intro + anchor, 1))

            rp = app / "RELEASE_MANIFEST.json"
            m = json.loads(rp.read_text())
            m["release"] = "16.1.91"
            m["release_label"] = "16.1.91 — notification readability and inventory date/workbook clarity"
            m["classification"] = "Presentation and workbook-instruction update derived from the exact reviewed 16.1.90 installer; no ERP business logic or data contract change."
            m["baseline"] = "Exact 16.1.90 canonical installer SHA-256 " + BASE_SHA
            m["frontend"]["notification_typography_16_1_91"] = "Role-independent notification title/message/metadata/button/toast scale with action/warning/danger/info severity accents; existing hooks and plain-text rendering preserved."
            m["frontend"]["inventory_date_help_16_1_91"] = "Count Date means the actual physical-count date (YYYY-MM-DD); posted ledger and SKU effective-date history remain controlled by authorized workflows."
            m["release_candidate_changes"].extend([
                "16.1.91: improve shared notification typography and action/warning/danger/info severity cues without changing role authorization or click handlers.",
                "16.1.91: clarify editable workbook Count Date and protect historical stock/SKU-effective-date records from direct Excel edits.",
                "16.1.91: plugin and frontend route/cache markers bumped; no schema, payroll, ledger, authorization or workflow transition changes."
            ])
            m["quality"] = {
                "release": "16.1.91",
                "local_validation_status": "Exact version is validated by the 16.1.91 GitHub workflows and attached build/QA artifact; this statement concerns source/package checks only.",
                "target_host_acceptance": "PENDING — no authenticated live WordPress/MariaDB run was performed by the build.",
                "production_certification": "PENDING — live integrity/ledger reconciliation, role-scoped acceptance, device PWA, verified backup restore and rollback are not performed by CI.",
                "no_unverified_live_acceptance_claims": True,
                "exact_16_1_91_github_actions_status": "See current exact-version workflow run."
            }
            m["release_acceptance"] = {
                "source_package_gate": "Exact-version CI status recorded by GitHub Actions; not a substitute for live acceptance.",
                "target_host": "PENDING — no authenticated live production execution performed.",
                "production_deployment": "NOT PERFORMED by build job.",
                "note": "Package/source checks do not certify live data integrity, ledger continuity, role acceptance, backup restoration or rollback."
            }
            m["corrective_release_16_1_91"] = {
                "scope": "Shared notification typography/severity accents and editable workbook date/history guidance.",
                "baseline_release": "16.1.90", "business_table_contract": "49 unchanged (29 Core/Operations + 16 Accounting/Procurement + 4 HR/Payroll)",
                "schema_changed": False, "database_version_changed": False, "inventory_ledger_changed": False,
                "sku_history_or_effective_dates_rewritten": False, "payroll_calculation_changed": False,
                "role_capability_ids_changed": False, "rest_business_logic_changed": False,
                "css_changed": True, "javascript_changed": False, "template_protection": "Both workbooks remain unprotected.",
                "live_production_deployment": False, "target_host_acceptance": "PENDING"
            }
            rp.write_text(json.dumps(m, ensure_ascii=False, indent=2) + "\n")

            files = sorted(p for p in app.rglob("*") if p.is_file())
            need(len(files) == 74, "package must have 74 files")
            with zipfile.ZipFile(OUT, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as out:
                for p in files:
                    info = zipfile.ZipInfo("stockland-erp-app/" + p.relative_to(app).as_posix(), (2026, 1, 1, 0, 0, 0))
                    info.compress_type = zipfile.ZIP_DEFLATED
                    info.external_attr = (0o100644 & 0xffff) << 16
                    out.writestr(info, p.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    digest = hashlib.sha256(OUT.read_bytes()).hexdigest()
    (REPO / (OUT.name + ".sha256")).write_text(digest + "  " + OUT.name + "\n")
    print("Built 16.1.91: " + str(OUT.stat().st_size) + " bytes, SHA-256 " + digest)

if __name__ == "__main__":
    main()
