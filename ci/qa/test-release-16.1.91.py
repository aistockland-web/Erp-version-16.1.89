#!/usr/bin/env python3
"""Independent package and source-contract checks for canonical ERP 16.1.91."""
import hashlib, json, re, zipfile, io
from pathlib import Path
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
BASE = REPO / "STOCKLAND_ERP_16.1.90_CANONICAL_WORDPRESS_UPDATE.zip"
OUT = REPO / "STOCKLAND_ERP_16.1.91_CANONICAL_WORDPRESS_UPDATE.zip"
SIDE = REPO / (OUT.name + ".sha256")

def need(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    if not ok:
        raise SystemExit(1)

need(BASE.is_file() and OUT.is_file() and SIDE.is_file(), "base, candidate and checksum sidecar exist")
need(hashlib.sha256(BASE.read_bytes()).hexdigest() == "6a7619c834fb37ef69024e68f15ae524efae39b5586b01441ab47270842588e5", "exact 16.1.90 baseline checksum")
raw = OUT.read_bytes()
sha = hashlib.sha256(raw).hexdigest()
need(sha in SIDE.read_text(encoding="utf-8"), "candidate checksum sidecar matches bytes")
with zipfile.ZipFile(BASE) as bz, zipfile.ZipFile(OUT) as z:
    need(bz.testzip() is None and z.testzip() is None, "baseline and candidate ZIP CRC")
    old_names, names = bz.namelist(), z.namelist()
    need(len(names) == 74 and len(set(names)) == 74, "74 unique plugin files")
    need(all(n.startswith("stockland-erp-app/") and ".." not in Path(n).parts for n in names), "safe single plugin root")
    need(set(names) == set(old_names), "no package paths added or removed")
    old = {n: bz.read(n) for n in old_names}
    new = {n: z.read(n) for n in names}
    changed = {n for n in old if old[n] != new[n]}
    allowed = {
        "stockland-erp-app/README.md", "stockland-erp-app/RELEASE_MANIFEST.json",
        "stockland-erp-app/assets/app.css", "stockland-erp-app/assets/app.min.css",
        "stockland-erp-app/assets/css/modules.json",
        *("stockland-erp-app/assets/css/" + name for name in json.loads(new["stockland-erp-app/assets/css/modules.json"])["modules"]),
        "stockland-erp-app/includes/class-stockland-frontend.php",
        "stockland-erp-app/stockland-inventory.php",
        "stockland-erp-app/templates/Stockland_Master_Import_Template.xlsx",
        "stockland-erp-app/templates/Stockland_Physical_Count_Template.xlsx",
    }
    need(changed == allowed and len(changed) == 15, "exact 15-path change allowlist, no runtime behavior files outside scope")
    need(sum(n.endswith(".php") for n in names) == 47, "47 runtime PHP files")
    app = "stockland-erp-app/"
    plugin = new[app + "stockland-inventory.php"].decode()
    front = new[app + "includes/class-stockland-frontend.php"].decode()
    manifest = json.loads(new[app + "RELEASE_MANIFEST.json"])
    modules = json.loads(new[app + "assets/css/modules.json"])
    need(manifest["release"] == modules["release"] == "16.1.91", "plugin, release and CSS module markers are 16.1.91")
    need("Version:           16.1.91" in plugin and "define( 'STOCKLAND_VERSION', '16.1.91' );" in plugin and "const ROUTE_VERSION      = '16.1.91';" in front, "plugin header/version/cache identity aligned")
    for schema in ["define( 'STOCKLAND_CORE_SCHEMA_VERSION', '16.1.1' );",
                   "define( 'STOCKLAND_ACCOUNTING_SCHEMA_VERSION', '16.1.15' );",
                   "define( 'STOCKLAND_HR_SCHEMA_VERSION', '16.1.3' );"]:
        need(schema in plugin, "database schema marker preserved: " + schema)
    for path in ["assets/app.js", "assets/pwa.js", "assets/bootstrap.min.css",
                 "includes/rest/class-rest-app.php", "includes/rest/class-rest-accounting.php",
                 "includes/rest/class-rest-hr.php", "includes/rest/class-rest-wholesale.php",
                 "includes/class-stockland-accounting.php", "includes/class-stockland-audit.php",
                 "includes/class-stockland-authority.php", "includes/class-stockland-db.php",
                 "includes/class-stockland-activator.php", "includes/class-stockland-payroll-statutory.php"]:
        need(new[app + path] == old[app + path], "preserved source byte-identically: " + path)

    module_names = modules["modules"]
    css_modules = [new[app + "assets/css/" + n].decode().strip() for n in module_names]
    expanded = "\n\n".join(css_modules) + "\n"
    need(expanded == new[app + "assets/app.css"].decode(), "app.css exactly rebuilds from six ordered source modules")
    normalize = lambda v: re.sub(r";}", "}", re.sub(r"\s*([{}:;,])\s*", r"\1", re.sub(r"\s+", " ", re.sub(r"/\*[\s\S]*?\*/", " ", v)).strip()))
    min_css = new[app + "assets/app.min.css"].decode().strip()
    need(normalize(expanded) == min_css, "app.min.css equals the canonical CSS normalization")
    need(min_css.startswith(old[app + "assets/app.min.css"].decode().strip()), "existing minified stylesheet preserved as prefix")
    for sel in [".sl-notification-card.action", ".sl-notification-card.success", ".sl-notification-card.warning", ".sl-notification-card.danger", ".sl-notification-card.info"]:
        need(sel in min_css, "semantic notification severity style exists: " + sel)
    for hook in ["data-review-notification", "data-review-notification-id", "data-review-purchase-order",
                 "data-review-purchase-order-id", "data-read-notification"]:
        need(hook in new[app + "assets/app.js"].decode(), "existing notification action hook retained: " + hook)

    physical = load_workbook(io.BytesIO(new[app + "templates/Stockland_Physical_Count_Template.xlsx"]))
    need(not physical["PHYSICAL COUNT"].protection.sheet and not physical["INSTRUCTIONS"].protection.sheet, "physical-count sheets remain editable")
    need("YYYY-MM-DD" in physical["PHYSICAL COUNT"]["A2"].value, "Count Date input format documented")
    need("date the physical count was actually taken" in physical["INSTRUCTIONS"]["A12"].value, "Count Date definition documented")
    need("historical SKU effective dates" in physical["INSTRUCTIONS"]["A13"].value, "immutable-history guardrail documented")
    master = load_workbook(io.BytesIO(new[app + "templates/Stockland_Master_Import_Template.xlsx"]))
    need(all(not ws.protection.sheet for ws in master.worksheets), "all seven master-import sheets remain editable")
    need("never creates physical stock" in master["PRODUCT MASTER"]["A2"].value.lower(), "product import remains quantity-free")
    need(master["IMPORT RULES"]["A9"].value == "SKU history / dates", "SKU date-history import rule added")
print("PASS exact 16.1.91 package contract")
print("SHA256=" + sha)
print("SIZE_BYTES=" + str(len(raw)))
